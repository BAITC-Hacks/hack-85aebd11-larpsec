import codecs
import re
from dataclasses import dataclass, field
from io import StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from pypdf.generic import ContentStream, StreamObject

from app.config import Settings
from app.errors import AppError
from app.models import Fragment

SUPPORTED = {".docx", ".pdf", ".xlsx", ".txt"}
CLAUSE = re.compile(r"^\s*(\d+(?:\.\d+)*)(?:[.)]|\s)")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class Parsed:
    fragments: list[Fragment] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    complete: bool = True


def word_text(paragraph) -> str:
    """Preserve Word's explicit separators without duplicating nested text boxes."""
    parts = []

    def visit(node):
        for child in node:
            if child.tag == W + "p":
                continue  # Nested paragraphs have their own locator in the caller.
            if child.tag == W + "t":
                parts.append(child.text or "")
            elif child.tag in {W + "br", W + "cr"}:
                parts.append("\n")
            elif child.tag == W + "tab":
                parts.append("\t")
            else:
                visit(child)

    visit(paragraph)
    return "".join(parts)


def pdf_has_images(page, reader) -> bool:
    """Inspect image/form resources and inline operators without decoding bitmap data."""
    visited = set()

    def visit(node):
        node = node.get_object()
        if id(node) in visited:
            return False
        visited.add(id(node))
        if node.get("/Subtype") == "/Image":
            return True
        stream = (
            ContentStream(node, reader) if isinstance(node, StreamObject) else node.get_contents()
        )
        if stream is not None:
            if len(stream.get_data()) > 20 * 1024**2:
                raise AppError("page_too_large", "Содержимое страницы PDF превышает лимит.", 413)
            if any(operator == b"INLINE IMAGE" for _, operator in stream.operations):
                return True
        resources = node.get("/Resources")
        if resources is None:
            return False
        objects = resources.get_object().get("/XObject")
        return objects is not None and any(visit(obj) for obj in objects.get_object().values())

    return visit(page)


def pdf_has_forms(page) -> bool:
    # pypdf layout extraction does not visit Form XObjects. Even a page with a
    # nonempty layout result may have its main text inside an embedded form.
    resources = page.get("/Resources")
    if resources is None:
        return False
    objects = resources.get_object().get("/XObject")
    return objects is not None and any(
        obj.get_object().get("/Subtype") == "/Form"
        for obj in objects.get_object().values()
    )


def check_office_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(x.file_size for x in entries) > 100 * 1024**2:
                raise AppError(
                    "archive_too_large", "Распакованный документ превышает 100 МБ или 10000 файлов."
                )
            if any(x.flag_bits & 1 for x in entries):
                raise AppError("encrypted_file", "Зашифрованные документы не поддерживаются.")
    except BadZipFile as exc:
        raise AppError(
            "invalid_document", "Файл не является корректным документом Office."
        ) from exc


def decode_plain_text(data: bytes) -> tuple[str, str | None]:
    # Do not reinterpret a renamed document or binary container as legacy text.
    binary_headers = (
        b"%PDF-",
        b"PK\x03\x04",
        b"PK\x05\x06",
        b"PK\x07\x08",
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
        b"\x89PNG\r\n\x1a\n",
        b"\xff\xd8\xff",
        b"GIF87a",
        b"GIF89a",
        b"\x1f\x8b",
        b"7z\xbc\xaf\x27\x1c",
        b"Rar!\x1a\x07",
        b"\x7fELF",
    )
    if data.lstrip(b" \t\r\n").startswith(binary_headers):
        raise AppError(
            "invalid_document",
            "Содержимое TXT не является обычным текстом. Загрузите документ с исходным расширением или сохраните его как текст UTF-8.",
        )
    warning = None
    encoding = None
    # UTF-32 LE shares its first two bytes with UTF-16 LE, so check it first.
    for markers, candidate in (
        ((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE), "utf-32"),
        ((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE), "utf-16"),
        ((codecs.BOM_UTF8,), "utf-8-sig"),
    ):
        if data.startswith(markers):
            encoding = candidate
            break
    try:
        if encoding:
            # A declared Unicode encoding must decode strictly. Falling back to
            # Windows-1251 here would silently disguise a damaged Unicode file.
            text = data.decode(encoding)
        else:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError as exc:
                # A valid non-ASCII UTF-8 prefix followed by damaged bytes is
                # ambiguous; require a clean export instead of inventing text.
                if any(value >= 128 for value in data[: exc.start]):
                    raise
                text = data.decode("cp1251")
                warning = (
                    "TXT прочитан как Windows-1251. Проверьте отображение букв; если оно неверное, сохраните исходный файл в UTF-8 и загрузите заново."
                )
    except UnicodeError as exc:
        raise AppError(
            "invalid_document",
            "Не удалось определить кодировку TXT или файл повреждён. Сохраните исходный текст в UTF-8 и загрузите заново.",
        ) from exc
    # Line breaks, tabs and form feeds are meaningful plain text separators.
    # Other C0/C1 controls (including NUL) indicate binary or damaged input.
    if re.search(r"[\x00-\x08\x0b\x0e-\x1f\x7f-\x9f]", text):
        raise AppError(
            "invalid_document",
            "TXT содержит бинарные данные или недопустимые управляющие символы. Сохраните файл как обычный текст UTF-8.",
        )
    return text, warning


def parse_document(path: Path, document_id: str, settings: Settings) -> Parsed:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise AppError("unsupported_format", "Поддерживаются DOCX, PDF, XLSX и TXT.", 415)
    result = Parsed()
    total_chars = 0

    def add(text: str, locator: str, **extra: object) -> None:
        nonlocal total_chars
        text = text.replace("\x00", "").strip()
        if not text:
            return
        if re.fullmatch(r"\d+(?:\.\d+)+\.?\s*[;:–—-]?", text):
            result.warnings.append(
                f"{locator}: пустой нумерованный пункт «{text}»; это не доказательство потери функции."
            )
        total_chars += len(text)
        if total_chars > settings.max_document_chars:
            raise AppError(
                "document_too_long",
                "В документе слишком много текста; разделите его на части.",
                413,
            )
        # Keep fragments bounded for extraction prompts, never discard the tail.
        for start in range(0, len(text), 6000):
            part = text[start : start + 6000]
            clause = CLAUSE.match(text)
            result.fragments.append(
                Fragment(
                    id=f"{document_id}:{len(result.fragments) + 1:05d}",
                    document_id=document_id,
                    text=part,
                    locator=locator + (f", часть {start // 6000 + 1}" if len(text) > 6000 else ""),
                    clause=clause.group(1) if clause else None,
                    **extra,
                )
            )

    try:
        if suffix in {".docx", ".xlsx"}:
            check_office_archive(path)
        if suffix == ".docx":
            doc = WordDocument(path)

            def add_paragraph(paragraph, locator):
                lines = word_text(paragraph).splitlines()
                for number, line in enumerate(lines, 1):
                    add(line, locator + (f", строка {number}" if len(lines) > 1 else ""))

            for index, paragraph in enumerate(doc.element.body.iter(W + "p"), 1):
                add_paragraph(paragraph, f"Абзац {index} (включая таблицы)")
            seen_parts = set()
            for section in doc.sections:
                for part in (section.header, section.footer):
                    if part.part.partname in seen_parts:
                        continue
                    seen_parts.add(part.part.partname)
                    for index, p in enumerate(part._element.iter(W + "p"), 1):
                        add_paragraph(
                            p,
                            f"Колонтитул {part.part.partname}, абзац {index}",
                        )
            if doc.element.body.xpath(".//w:drawing | .//w:pict | .//w:object"):
                result.complete = False
                result.warnings.append(
                    "Есть рисунки или вложенные объекты: их содержимое не распознано; требуется проверка/OCR."
                )
            if doc.element.body.xpath(".//w:ins | .//w:del"):
                result.complete = False
                result.warnings.append(
                    "Обнаружены исправления Word: загрузите редакцию с принятыми изменениями."
                )
        elif suffix == ".pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise AppError("encrypted_file", "Снимите пароль с PDF перед загрузкой.")
            if len(reader.pages) > 300:
                raise AppError("too_many_pages", "Допускается не более 300 страниц PDF.", 413)
            for page_number, page in enumerate(reader.pages, 1):
                content = page.get_contents()
                if content is not None and len(content.get_data()) > 20 * 1024**2:
                    raise AppError(
                        "page_too_large", "Содержимое страницы PDF превышает лимит.", 413
                    )
                text = ""
                if content is not None:
                    try:
                        # Group words by their page positions, including rotated
                        # text; content-stream order can put every word on a line.
                        if not pdf_has_forms(page):
                            text = page.extract_text(
                                extraction_mode="layout",
                                layout_mode_space_vertically=False,
                                layout_mode_strip_rotated=False,
                            ) or ""
                    except Exception:
                        # Some PDF fonts/layouts cannot be handled geometrically.
                        # The plain extractor may still recover useful evidence.
                        text = ""
                    if not text.strip():
                        text = page.extract_text() or ""
                        if text.strip():
                            result.complete = False
                            result.warnings.append(
                                f"Страница {page_number}: текст извлечён без сохранения расположения строк; проверьте порядок слов по оригиналу."
                            )
                if pdf_has_images(page, reader):
                    result.complete = False
                    result.warnings.append(
                        f"Страница {page_number}: есть изображение; часть содержимого может требовать OCR даже при наличии текстового слоя."
                    )
                if len(text.strip()) < 10:
                    result.complete = False
                    result.warnings.append(
                        f"Страница {page_number}: недостаточно текста; возможен скан, требуется OCR."
                    )
                for index, line in enumerate(text.splitlines(), 1):
                    add(line, f"Страница {page_number}, строка {index}", page=page_number)
        elif suffix == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in book.worksheets:
                    if (sheet.max_row or 0) > 50000 or (sheet.max_column or 0) > 200:
                        raise AppError(
                            "sheet_too_large", "Лист превышает 50000 строк или 200 столбцов.", 413
                        )
                    count = 0
                    has_formulas = False
                    for row in sheet.iter_rows():
                        count += len(row)
                        if count > 200000:
                            raise AppError("sheet_too_large", "Лист превышает 200000 ячеек.", 413)
                        cells = [c for c in row if c.value is not None]
                        if not cells:
                            continue
                        has_formulas |= any(c.data_type == "f" for c in cells)
                        cell_range = f"{get_column_letter(cells[0].column)}{cells[0].row}:{get_column_letter(cells[-1].column)}{cells[-1].row}"
                        add(
                            " | ".join(f"{c.coordinate}: {c.value}" for c in cells),
                            f"Лист «{sheet.title}», {cell_range}",
                            sheet=sheet.title,
                            cell_range=cell_range,
                        )
                    if has_formulas:
                        result.complete = False
                        result.warnings.append(
                            f"Лист «{sheet.title}»: формулы сохранены как текст, вычисления не выполнялись."
                        )
                with ZipFile(path) as archive:
                    if any(n.startswith(("xl/drawings/", "xl/media/")) for n in archive.namelist()):
                        result.complete = False
                        result.warnings.append(
                            "В Excel есть графические объекты, их содержимое требует отдельной проверки."
                        )
            finally:
                book.close()
        elif suffix == ".txt":
            byte_limit = settings.max_upload_mb * 1024**2
            with path.open("rb") as source:
                data = source.read(byte_limit + 1)
            if len(data) > byte_limit:
                raise AppError("upload_too_large", "Файл превышает допустимый размер.", 413)
            text, warning = decode_plain_text(data)
            if warning:
                result.warnings.append(warning)
            # Preserve physical source line numbers, including skipped empty
            # lines, so citations point back to the original TXT accurately.
            with StringIO(text, newline=None) as lines:
                for index, line in enumerate(lines, 1):
                    add(line, f"Строка {index}")
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "invalid_document",
            "Не удалось прочитать документ: проверьте формат и целостность файла.",
        ) from exc
    if not result.fragments:
        if suffix == ".txt":
            raise AppError("no_text", "Файл TXT пуст или содержит только пробелы и переносы строк.")
        raise AppError("no_text", "Текст не найден. Для сканов сначала выполните OCR.")
    return result
