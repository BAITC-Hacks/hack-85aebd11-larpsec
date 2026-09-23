from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.demo import extract_demo
from app.errors import AppError
from app.parsing import parse_document


def test_docx_keeps_table_order_and_clauses(tmp_path, settings):
    doc = Document()
    doc.add_paragraph("1.1. Первая функция")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "1.2. Функция в таблице"
    doc.add_paragraph("1.3. Последняя функция")
    path = tmp_path / "table.docx"
    doc.save(path)
    parsed = parse_document(path, "doc", settings)
    assert [f.clause for f in parsed.fragments] == ["1.1", "1.2", "1.3"]
    assert len({f.id for f in parsed.fragments}) == 3
    assert all(f.page is None for f in parsed.fragments)  # DOCX has no reliable page layout.


def test_xlsx_preserves_sheet_cells_and_warns_on_formulas(tmp_path, settings):
    book = Workbook()
    book.active.title = "Функции"
    book.active.append(["Департамент", "Проверяет отчётность"])
    book.active.append(["=1+1", "Формула"])
    path = tmp_path / "functions.xlsx"
    book.save(path)
    parsed = parse_document(path, "doc", settings)
    assert parsed.fragments[0].sheet == "Функции" and parsed.fragments[0].cell_range == "A1:B1"
    assert parsed.fragments[0].text == "Департамент | Проверяет отчётность"
    assert not parsed.complete and parsed.warnings


def test_demo_xlsx_and_docx_extract_same_functions(demo_files, settings):
    for side in ("before", "after"):
        word = extract_demo(parse_document(demo_files / f"{side}.docx", "word", settings).fragments)
        excel = extract_demo(
            parse_document(demo_files / f"{side}.xlsx", "excel", settings).fragments
        )
        assert [(f.owner, f.action) for f in word.functions] == [
            (f.owner, f.action) for f in excel.functions
        ]


def test_pdf_text_and_mixed_scan_warning(tmp_path, settings):
    path = tmp_path / "mixed.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 700, "1.1. Department audits financial reports.")
    pdf.showPage()
    pdf.rect(0, 0, 100, 100)
    pdf.showPage()
    pdf.save()
    parsed = parse_document(path, "doc", settings)
    assert parsed.fragments[0].page == 1 and parsed.fragments[0].clause == "1.1"
    assert not parsed.complete and any("2" in w for w in parsed.warnings)


def test_empty_pdf_requires_ocr(tmp_path, settings):
    path = tmp_path / "scan.pdf"
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.write(path)
    with pytest.raises(AppError, match="OCR"):
        parse_document(path, "doc", settings)


def test_encrypted_pdf_rejected(tmp_path, settings):
    path = tmp_path / "secret.pdf"
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.encrypt("secret")
    pdf.write(path)
    with pytest.raises(AppError) as error:
        parse_document(path, "doc", settings)
    assert error.value.code == "encrypted_file"


def test_long_text_is_bounded_without_silent_truncation(tmp_path, settings):
    doc = Document()
    doc.add_paragraph("а" * 7000)
    path = tmp_path / "long.docx"
    doc.save(path)
    parsed = parse_document(path, "doc", settings)
    assert len(parsed.fragments) == 2 and sum(len(f.text) for f in parsed.fragments) == 7000
    settings.max_document_chars = 1000
    with pytest.raises(AppError) as error:
        parse_document(path, "doc", settings)
    assert error.value.code == "document_too_long"


def test_malformed_input_is_safe(tmp_path, settings):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a PDF")
    with pytest.raises(AppError) as error:
        parse_document(path, "doc", settings)
    assert error.value.code == "invalid_document"


def test_word_image_marks_coverage_incomplete(tmp_path, settings):
    from PIL import Image

    stream = BytesIO()
    Image.new("RGB", (20, 20), "white").save(stream, format="PNG")
    stream.seek(0)
    doc = Document()
    doc.add_paragraph("1.1. Есть текст и изображение структуры.")
    doc.add_picture(stream)
    path = tmp_path / "image.docx"
    doc.save(path)
    parsed = parse_document(path, "doc", settings)
    assert not parsed.complete and parsed.warnings


def test_word_manual_line_break_and_tab_do_not_merge_owner_and_function(tmp_path, settings):
    doc = Document()
    paragraph = doc.add_paragraph("Подразделение: Служба аудита")
    paragraph.add_run().add_break()
    paragraph.add_run("1.1. Проверяет")
    paragraph.add_run().add_tab()
    paragraph.add_run("сохранность активов.")
    path = tmp_path / "linebreak.docx"
    doc.save(path)
    parsed = parse_document(path, "word", settings)
    assert [f.text for f in parsed.fragments] == [
        "Подразделение: Служба аудита",
        "1.1. Проверяет\tсохранность активов.",
    ]
    draft = extract_demo(parsed.fragments)
    assert draft.functions[0].owner == "Служба аудита"
    assert "Проверяет сохранность" in draft.functions[0].action
    assert "Проверяет\tсохранность" in draft.functions[0].evidence[0].quote


@pytest.mark.parametrize("inline", [False, True])
def test_pdf_scan_with_readable_header_is_still_incomplete(tmp_path, settings, inline):
    from PIL import Image
    from reportlab.lib.utils import ImageReader

    path = tmp_path / "scan-with-header.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 780, "Internal audit regulation: readable header")
    bitmap = Image.new("RGB", (100, 100), "white")
    if inline:
        pdf.drawInlineImage(bitmap, 50, 100, width=400, height=600)
    else:
        pdf.drawImage(ImageReader(bitmap), 50, 100, width=400, height=600)
    pdf.showPage()
    pdf.save()
    parsed = parse_document(path, "scan", settings)
    assert parsed.fragments and not parsed.complete
    assert any("изображение" in warning for warning in parsed.warnings)
