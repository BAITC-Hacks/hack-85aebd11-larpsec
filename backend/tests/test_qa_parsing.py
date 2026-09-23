import codecs
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from pypdf import PageObject, PdfReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from app.demo import extract_demo
from app.errors import AppError
from app.parsing import clause_number, parse_document, pdf_paragraphs
from tests.conftest import analyze, upload


def test_xlsx_header_table_emits_owned_functions_without_cell_addresses(tmp_path, settings):
    book = Workbook()
    sheet = book.active
    sheet.title = "Обязанности"
    sheet.append(["№", "Подразделение", "Функция", "Область"])
    sheet.append(["1.1", "Отдел аудита", "Проверяет права доступа.", "Филиалы"])
    sheet.append(["1.2", None, "Готовит отчёт.", "Филиалы"])
    sheet.append(["2.1", "Отдел контроля", "Проверяет сохранность активов.", None])
    sheet.append(["1.3", "Отдел аудита", "Организует обучение.", None])
    path = tmp_path / "table.xlsx"
    book.save(path)

    parsed = parse_document(path, "sheet", settings)
    draft = extract_demo(parsed.fragments)
    assert parsed.complete
    assert [f.owner for f in draft.functions] == [
        "Отдел аудита",
        "Отдел аудита",
        "Отдел контроля",
        "Отдел аудита",
    ]
    assert [f.clause for f in parsed.fragments if f.clause] == ["1.1", "1.2", "2.1", "1.3"]
    first = next(f for f in parsed.fragments if f.clause == "1.1")
    assert first.text == "1.1. Проверяет права доступа.; Область: Филиалы"
    assert first.cell_range == "A2:D2" and first.sheet == "Обязанности"
    assert "A2:D2" in first.locator
    assert all("A2:" not in f.text and "C2:" not in f.text for f in parsed.fragments)
    assert draft.functions[0].evidence[0].quote == first.text


@pytest.mark.parametrize(("heading", "marker"), [("Должность", "Должность"), ("Группа", "Группа")])
def test_xlsx_owner_kind_and_norm_are_preserved(tmp_path, settings, heading, marker):
    book = Workbook()
    book.active.append(["Пункт", heading, "Обязанность", "Тип нормы"])
    book.active.append([3, "Аудитор", "утверждать платежи", "Запрет"])
    path = tmp_path / "typed.xlsx"
    book.save(path)
    parsed = parse_document(path, "typed", settings)
    assert any(f.text == f"{marker}: Аудитор" for f in parsed.fragments)
    function = next(f for f in parsed.fragments if f.clause == "3")
    assert function.text == "3. Запрещено: утверждать платежи"
    assert function.cell_range == "A2:D2"


@pytest.mark.parametrize(
    "header", [None, ["№", "Функция"], ["Подразделение", "Должность", "Функция"]]
)
def test_xlsx_ambiguous_structure_keeps_values_and_marks_incomplete(tmp_path, settings, header):
    book = Workbook()
    if header:
        book.active.append(header)
    if header == ["№", "Функция"]:
        book.active.append([1, "Проверяет отчётность"])
    else:
        book.active.append(["Отдел аудита", "Аудитор", "Проверяет отчётность"])
    path = tmp_path / "unknown.xlsx"
    book.save(path)
    parsed = parse_document(path, "unknown", settings)
    assert not parsed.complete and parsed.warnings
    assert any("Проверяет отчётность" in f.text for f in parsed.fragments)


def test_pdf_wrapped_numbered_paragraphs_keep_full_action_and_line_locator(tmp_path, settings):
    path = tmp_path / "wrapped.pdf"
    pdf = canvas.Canvas(str(path))
    text = pdf.beginText(50, 760)
    for line in [
        "Audit policy",
        "1.1. Checks   internal",
        "    controls and prepares",
        "reports.",
        "1.2. Verifies access rights.",
    ]:
        text.textLine(line)
    pdf.drawText(text)
    pdf.save()
    parsed = parse_document(path, "pdf", settings)
    assert [f.text for f in parsed.fragments] == [
        "Audit policy",
        "1.1. Checks internal controls and prepares reports.",
        "1.2. Verifies access rights.",
    ]
    assert parsed.fragments[1].clause == "1.1" and parsed.fragments[1].page == 1
    assert parsed.fragments[1].locator == "Страница 1, строки 2–4"


@pytest.mark.parametrize(
    "heading",
    [
        "Переименована из: Служба контроля",
        "Переименовано из: Управление контроля",
        "Переименованы из: Группа контроля",
        "Прежнее наименование: Отдел контроля",
        "Ранее: Отдел контроля",
        "Положение об отделе контроля",
        "Функции отдела контроля",
        "Обязанности отдела контроля",
        "ОТВЕТСТВЕННОСТЬ",
    ],
)
def test_pdf_structural_headings_are_not_swallowed_by_previous_function(heading):
    text = f"1.1. Проверяет отчётность.\n{heading}\n2.1. Готовит отчёт."
    fragments = list(pdf_paragraphs(text))
    assert [part[0] for part in fragments] == [
        "1.1. Проверяет отчётность.",
        heading,
        "2.1. Готовит отчёт.",
    ]


@pytest.mark.parametrize("marker", ["-", "–", "—", "•", "·", "*"])
def test_pdf_bullets_start_a_new_function_and_keep_their_wrapped_tail(marker):
    text = f"1.1. Проверяет отчётность.\n{marker} Готовит отчёт\nпо результатам проверки."
    assert [part[0] for part in pdf_paragraphs(text)] == [
        "1.1. Проверяет отчётность.",
        f"{marker} Готовит отчёт по результатам проверки.",
    ]


def test_pdf_vertical_gap_separates_paragraph_without_splitting_indented_wrap(tmp_path, settings):
    path = tmp_path / "spaced.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 760, "1.1. Checks accounts")
    pdf.drawString(68, 745, "of local branches.")
    pdf.drawString(160, 675, "Separate unnumbered note.")
    pdf.save()
    parsed = parse_document(path, "spaced", settings)
    assert [f.text for f in parsed.fragments] == [
        "1.1. Checks accounts of local branches.",
        "Separate unnumbered note.",
    ]
    assert [f.clause for f in parsed.fragments] == ["1.1", None]


def test_russian_pdf_changed_wrapped_tail_is_not_reported_as_preserved(client, tmp_path):
    font = next(
        (
            p
            for p in [
                Path("/usr/share/fonts/TTF/DejaVuSans.ttf"),
                Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            ]
            if p.is_file()
        ),
        None,
    )
    if font is None:
        pytest.skip("DejaVu Sans is needed to generate the Cyrillic PDF fixture")
    pdfmetrics.registerFont(TTFont("QADejaVu", str(font)))
    cid = client.post(
        "/api/v1/comparisons", json={"before_complete": True, "after_complete": True}
    ).json()["id"]
    for side, tail in [("before", "филиалов."), ("after", "головного офиса.")]:
        path = tmp_path / f"{side}.pdf"
        pdf = canvas.Canvas(str(path))
        pdf.setFont("QADejaVu", 11)
        text = pdf.beginText(50, 760)
        text.setLeading(14)
        for line in ["Подразделение: Отдел аудита", "1.1. Проверяет отчётность", tail]:
            text.textLine(line)
        pdf.drawText(text)
        pdf.save()
        response = upload(client, cid, path, side)
        assert response.status_code == 201, response.text
    result = analyze(client, cid)
    before = next(f for f in result["functions"] if f["side"] == "before")
    after = next(f for f in result["functions"] if f["side"] == "after")
    assert before["action"] == "Проверяет отчётность филиалов"
    assert after["action"] == "Проверяет отчётность головного офиса"
    assert not any(change["status"] == "preserved" for change in result["function_changes"])
    assert before["evidence"][0]["quote"] == "1.1. Проверяет отчётность филиалов."


@pytest.mark.parametrize(
    "text",
    [
        "25.06.2021 Утверждено",
        "01.02.23. Утверждено",
        "2021 год",
        "2024. Отчёт",
        "10 сотрудников",
        "3.14 составляет значение",
        "1.5 кг",
        "12.5 процентов",
        "100. руб.",
        "25. процентов",
        "1.000.000 тенге",
    ],
)
def test_numbers_in_prose_are_not_clause_references(text):
    assert clause_number(text) is None


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("1.1. Проверяет отчёты", "1.1"),
        ("1) Проверяет отчёты", "1"),
        ("12. Проверяет отчёты", "12"),
        ("5.9.14. Проверяет отчёты", "5.9.14"),
        ("1.2 Проводит проверки", "1.2"),
        ("5.5.3 готовит отчёты", "5.5.3"),
    ],
)
def test_explicit_clause_markers_are_recognized(text, expected):
    assert clause_number(text) == expected


def element(name, **values):
    item = OxmlElement(f"w:{name}")
    for key, value in values.items():
        item.set(qn(f"w:{key}"), str(value))
    return item


def numbering_definition(doc, num_id=99, start=1):
    root = doc.part.numbering_part.element
    abstract = element("abstractNum", abstractNumId=num_id)
    for level, template in [(0, "%1."), (1, "%1.%2.")]:
        item = element("lvl", ilvl=level)
        item.extend(
            [
                element("start", val=1),
                element("numFmt", val="decimal"),
                element("lvlText", val=template),
            ]
        )
        abstract.append(item)
    root.append(abstract)
    number = element("num", numId=num_id)
    number.append(element("abstractNumId", val=num_id))
    if start != 1:
        override = element("lvlOverride", ilvl=0)
        override.append(element("startOverride", val=start))
        number.append(override)
    root.append(number)


def assign_numbering(properties, num_id, level=0):
    numbering = element("numPr")
    numbering.extend([element("ilvl", val=level), element("numId", val=num_id)])
    properties.append(numbering)


def test_docx_real_numbering_multilevel_restart_and_override(tmp_path, settings):
    doc = Document()
    numbering_definition(doc, start=4)
    for text, level in [
        ("Контроль", 0),
        ("Проверяет отчёты", 1),
        ("Проверяет активы", 1),
        ("Аудит", 0),
        ("Готовит отчёты", 1),
    ]:
        paragraph = doc.add_paragraph(text)
        assign_numbering(paragraph._p.get_or_add_pPr(), 99, level)
    path = tmp_path / "numbering.docx"
    doc.save(path)
    parsed = parse_document(path, "numbered", settings)
    assert [f.clause for f in parsed.fragments] == ["4", "4.1", "4.2", "5", "5.1"]
    assert parsed.fragments[1].text == "4.1. Проверяет отчёты"
    assert parsed.complete


def test_docx_numbering_is_inherited_from_style_chain(tmp_path, settings):
    doc = Document()
    numbering_definition(doc)
    parent = doc.styles.add_style("QA Parent List", WD_STYLE_TYPE.PARAGRAPH)
    assign_numbering(parent.element.get_or_add_pPr(), 99, 1)
    child = doc.styles.add_style("QA Child List", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    doc.add_paragraph("Проверяет отчёты", style=child)
    doc.add_paragraph("Готовит заключение", style=child)
    path = tmp_path / "style-numbering.docx"
    doc.save(path)
    parsed = parse_document(path, "styled", settings)
    assert [f.clause for f in parsed.fragments] == ["1.1", "1.2"]
    assert parsed.fragments[0].text == "1.1. Проверяет отчёты"


def test_docx_builtin_numbered_style_and_unknown_definition(tmp_path, settings):
    doc = Document()
    doc.add_paragraph("Первый пункт", style="List Number")
    doc.add_paragraph("Второй пункт", style="List Number")
    broken = doc.add_paragraph("Текст без доступного определения нумерации")
    assign_numbering(broken._p.get_or_add_pPr(), 999)
    path = tmp_path / "list.docx"
    doc.save(path)
    parsed = parse_document(path, "list", settings)
    assert [f.clause for f in parsed.fragments] == ["1", "2", None]
    assert parsed.fragments[2].text == "Текст без доступного определения нумерации"
    assert not parsed.complete and any("нумерацию Word" in w for w in parsed.warnings)


def test_docx_style_level_is_resolved_from_numbering_definition(tmp_path, settings):
    doc = Document()
    numbering_definition(doc)
    style = doc.styles.add_style("QA Outline Level", WD_STYLE_TYPE.PARAGRAPH)
    num_pr = element("numPr")
    num_pr.append(element("numId", val=99))  # Word may omit ilvl from the style.
    style.element.get_or_add_pPr().append(num_pr)
    abstract = next(
        n
        for n in doc.part.numbering_part.element.findall(qn("w:abstractNum"))
        if n.get(qn("w:abstractNumId")) == "99"
    )
    abstract.findall(qn("w:lvl"))[1].append(element("pStyle", val=style.style_id))
    doc.add_paragraph("Проверяет отчёты", style=style)
    path = tmp_path / "outline-style.docx"
    doc.save(path)
    parsed = parse_document(path, "outline", settings)
    assert parsed.fragments[0].text == "1.1. Проверяет отчёты"
    assert parsed.fragments[0].clause == "1.1"


def test_docx_headers_and_footers_use_friendly_section_locators(tmp_path, settings):
    doc = Document()
    doc.add_paragraph("Основной текст")
    doc.sections[0].header.paragraphs[0].text = "Заголовок первого раздела"
    doc.sections[0].footer.paragraphs[0].text = "Подвал первого раздела"
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "Заголовок второго раздела"
    path = tmp_path / "headers.docx"
    doc.save(path)
    parsed = parse_document(path, "headers", settings)
    locators = {f.text: f.locator for f in parsed.fragments}
    assert locators["Заголовок первого раздела"] == "Верхний колонтитул, раздел 1, абзац 1"
    assert locators["Подвал первого раздела"] == "Нижний колонтитул, раздел 1, абзац 1"
    assert locators["Заголовок второго раздела"] == "Верхний колонтитул, раздел 2, абзац 1"
    assert not any(".xml" in locator or "/word" in locator for locator in locators.values())


def test_existing_public_docx_without_numbering_part_still_parses_and_uploads(settings, client):
    path = (
        Path(__file__).resolve().parents[2]
        / "frontend"
        / "public"
        / "examples"
        / "audit-example.docx"
    )
    parsed = parse_document(path, "public-example", settings)
    assert parsed.fragments
    assert any(f.clause == "5.5.3" for f in parsed.fragments)
    cid = client.post("/api/v1/comparisons", json={}).json()["id"]
    response = upload(client, cid, path, "before")
    assert response.status_code == 201, response.text
    assert response.json()["fragment_count"] == len(parsed.fragments)


def test_extension_only_docx_name_can_be_parsed_and_uploaded(tmp_path, settings, client):
    doc = Document()
    doc.add_paragraph("Подразделение: Аудит")
    doc.add_paragraph("1.1. Проверяет отчёты.")
    path = tmp_path / ".docx"
    doc.save(path)
    assert parse_document(path, "extension", settings).fragments[1].clause == "1.1"
    cid = client.post("/api/v1/comparisons", json={}).json()["id"]
    response = client.post(
        f"/api/v1/comparisons/{cid}/documents?side=before",
        files={"file": (".docx", path.read_bytes())},
    )
    assert response.status_code == 201, response.text
    assert response.json()["filename"] == ".docx" and response.json()["format"] == "docx"


@pytest.mark.parametrize("filename", ["broken.docx", ".docx", "broken.xlsx", "broken.pdf"])
def test_malformed_upload_keeps_comparison_and_storage_unchanged(client, filename):
    cid = client.post("/api/v1/comparisons", json={}).json()["id"]
    original = client.get(f"/api/v1/comparisons/{cid}").json()
    response = client.post(
        f"/api/v1/comparisons/{cid}/documents?side=before",
        files={"file": (filename, b"not a document")},
    )
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "invalid_document"
    assert client.get(f"/api/v1/comparisons/{cid}").json() == original
    assert list(client.app.state.store.uploads.iterdir()) == []


@pytest.mark.parametrize(
    ("encoding", "prefix"),
    [
        ("utf-8", b""),
        ("utf-8-sig", b""),
        ("utf-16-le", codecs.BOM_UTF16_LE),
        ("utf-16-be", codecs.BOM_UTF16_BE),
        ("utf-32-le", codecs.BOM_UTF32_LE),
        ("utf-32-be", codecs.BOM_UTF32_BE),
        ("cp1251", b""),
    ],
)
def test_txt_encodings_keep_source_lines_and_strict_clauses(tmp_path, settings, encoding, prefix):
    text = "Подразделение: Аудит\r\n\r\n1.2 Проверяет отчётность.\r2025 год\n"
    path = tmp_path / ".txt"
    path.write_bytes(prefix + text.encode(encoding))
    parsed = parse_document(path, "txt", settings)
    assert [f.text for f in parsed.fragments] == [
        "Подразделение: Аудит",
        "1.2 Проверяет отчётность.",
        "2025 год",
    ]
    assert [f.locator for f in parsed.fragments] == ["Строка 1", "Строка 3", "Строка 4"]
    assert [f.clause for f in parsed.fragments] == [None, "1.2", None]
    assert bool(parsed.warnings) == (encoding == "cp1251")
    assert parsed.complete


@pytest.mark.parametrize(
    "data",
    [
        b" \n%PDF-1.4\n",
        b"PK\x03\x04document",
        b"plain\x00binary",
        b"plain\x7ftext",
        codecs.BOM_UTF8 + b"damaged\xff",
        codecs.BOM_UTF16_LE + b"\x41",
        codecs.BOM_UTF32_BE + b"\x00\x00\x00",
        "Текст".encode() + b"\xff",
    ],
)
def test_txt_rejects_binary_and_damaged_declared_unicode(tmp_path, settings, data):
    path = tmp_path / "invalid.txt"
    path.write_bytes(data)
    with pytest.raises(AppError) as error:
        parse_document(path, "txt", settings)
    assert error.value.code == "invalid_document"


def test_txt_empty_character_and_upload_limits(tmp_path, settings):
    path = tmp_path / "limits.txt"
    path.write_text(" \n\t\r\n", encoding="utf-8")
    with pytest.raises(AppError) as error:
        parse_document(path, "txt", settings)
    assert error.value.code == "no_text"

    text = "1.1. " + "Проверяет отчёты. " * 750
    path.write_text(text, encoding="utf-8")
    parsed = parse_document(path, "txt", settings)
    assert "".join(f.text for f in parsed.fragments) == text.strip()
    assert all(len(f.text) <= 6000 and f.clause == "1.1" for f in parsed.fragments)
    assert all(f.locator.startswith("Строка 1, часть ") for f in parsed.fragments)
    with pytest.raises(AppError) as error:
        parse_document(path, "txt", settings.model_copy(update={"max_document_chars": 1000}))
    assert error.value.code == "document_too_long" and error.value.status == 413

    path.write_bytes(b"x" * (1024**2 + 1))
    with pytest.raises(AppError) as error:
        parse_document(path, "txt", settings.model_copy(update={"max_upload_mb": 1}))
    assert error.value.code == "upload_too_large" and error.value.status == 413


def test_pdf_geometric_word_order_keeps_wrapped_paragraph_and_vertical_gap(tmp_path, settings):
    path = tmp_path / "word-positions.pdf"
    pdf = canvas.Canvas(str(path))
    words = []
    position = 50
    for word in "1.1. Checks branch accounts".split():
        words.append((position, word))
        position += pdfmetrics.stringWidth(word + " ", "Helvetica", 12)
    for position, word in reversed(words):
        pdf.drawString(position, 760, word)
    pdf.drawString(65, 745, "and prepares reports.")
    pdf.drawString(100, 675, "Separate unnumbered note.")
    pdf.save()
    plain = PdfReader(path).pages[0].extract_text()
    assert plain.index("accounts") < plain.index("Checks")
    parsed = parse_document(path, "pdf", settings)
    assert [f.text for f in parsed.fragments] == [
        "1.1. Checks branch accounts and prepares reports.",
        "Separate unnumbered note.",
    ]
    assert parsed.complete
    assert parsed.fragments[0].clause == "1.1"


def test_pdf_form_text_is_retained_even_when_page_has_regular_header(tmp_path, settings):
    path = tmp_path / "form.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 760, "Audit policy header")
    pdf.beginForm("embedded")
    pdf.drawString(50, 720, "1.1. Checks branch accounts")
    pdf.drawString(65, 705, "and prepares reports.")
    pdf.endForm()
    pdf.doForm("embedded")
    pdf.save()
    parsed = parse_document(path, "pdf", settings)
    assert [f.text for f in parsed.fragments] == [
        "Audit policy header",
        "1.1. Checks branch accounts and prepares reports.",
    ]
    assert not parsed.complete
    assert any("без сохранения расположения" in warning for warning in parsed.warnings)


def test_pdf_rotated_text_is_not_silently_stripped(tmp_path, settings):
    path = tmp_path / "rotated.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 760, "Audit policy header")
    pdf.saveState()
    pdf.translate(50, 500)
    pdf.rotate(90)
    pdf.drawString(0, 0, "Checks branch accounts.")
    pdf.restoreState()
    pdf.save()
    parsed = parse_document(path, "pdf", settings)
    assert "Checks branch accounts." in " ".join(f.text for f in parsed.fragments)


@pytest.mark.parametrize("failure", ["exception", "omitted_tail"])
def test_pdf_layout_fallback_preserves_evidence_and_warns(tmp_path, settings, monkeypatch, failure):
    path = tmp_path / "fallback.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 760, "1.1. Checks branch accounts and prepares reports.")
    pdf.save()
    original = PageObject.extract_text

    def extract(page, *args, **kwargs):
        if kwargs.get("extraction_mode") == "layout":
            if failure == "exception":
                raise ValueError("Unsupported geometry")
            return "1.1. Checks branch accounts."
        return original(page, *args, **kwargs)

    monkeypatch.setattr(PageObject, "extract_text", extract)
    parsed = parse_document(path, "pdf", settings)
    assert parsed.fragments[0].text == "1.1. Checks branch accounts and prepares reports."
    assert not parsed.complete and parsed.warnings


def test_xlsx_sparse_values_and_extra_columns_keep_content_and_cell_locators(tmp_path, settings):
    book = Workbook()
    sheet = book.active
    sheet["D2"] = "A standalone note after empty cells."
    sheet.append(["№", "Подразделение", "Функция", None, "Примечание"])
    sheet.append(["1.1", "Отдел аудита", "Проверяет отчёты.", None, "Только филиалы."])
    path = tmp_path / "sparse.xlsx"
    book.save(path)
    parsed = parse_document(path, "sheet", settings)
    note = next(f for f in parsed.fragments if f.text == "A standalone note after empty cells.")
    assert note.cell_range == "D2:D2" and "D2:D2" in note.locator
    function = next(f for f in parsed.fragments if f.clause == "1.1")
    assert function.text == "1.1. Проверяет отчёты."
    assert function.cell_range == "A4:E4"
    extra = next(f for f in parsed.fragments if f.text == "Примечание: Только филиалы.")
    assert extra.cell_range == "A4:E4"
    assert all("A4:" not in f.text and "D2:" not in f.text for f in parsed.fragments)
