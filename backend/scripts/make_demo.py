"""Generate clearly labeled synthetic input documents; never changes organizer files."""

import argparse
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

SCENARIO = Path(__file__).resolve().parents[2] / "examples" / "scenario.json"


def make_demo(output: Path, scenario_path: Path = SCENARIO) -> dict:
    scenario = json.loads(scenario_path.read_text())
    output.mkdir(parents=True, exist_ok=True)
    for side in ("before", "after"):
        doc = Document()
        for paragraph in scenario[side]:
            doc.add_paragraph(paragraph)
        doc.save(output / f"{side}.docx")
        book = Workbook()
        sheet = book.active
        sheet.title = "Учебные функции"
        for paragraph in scenario[side]:
            sheet.append([paragraph])
        book.save(output / f"{side}.xlsx")
    (output / "expected.json").write_text(
        json.dumps(scenario["expected"], ensure_ascii=False, indent=2) + "\n"
    )
    return scenario


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("data/demo"))
    args = parser.parse_args()
    make_demo(args.output)
    print(f"Учебные документы созданы: {args.output.resolve()}")
